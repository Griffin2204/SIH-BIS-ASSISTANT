import os
from typing import Dict, Any, Optional, List
from pypdf import PdfReader
import docx
from services.text_cleaner import clean_text
from services.chunker import chunk_text, DocumentChunk
from services.embedding_service import EmbeddingService, DEFAULT_MODEL_NAME
from services.vector_store import VectorStoreService
from services.security_utils import MAX_DOCUMENT_CHARS

class DocumentExtractionResult:
    def __init__(
        self,
        success: bool,
        raw_text: str = "",
        cleaned_text: str = "",
        chunks: Optional[List[DocumentChunk]] = None,
        file_type: str = "",
        text_length: int = 0,
        cleaned_text_length: int = 0,
        chunk_count: int = 0,
        pages: Optional[int] = None,
        embedding_model: str = DEFAULT_MODEL_NAME,
        embedding_dimension: int = 384,
        vector_store_status: str = "failed",
        error: Optional[str] = None
    ):
        self.success = success
        self.raw_text = raw_text
        self.cleaned_text = cleaned_text
        self.chunks = chunks or []
        self.file_type = file_type
        self.text_length = text_length
        self.cleaned_text_length = cleaned_text_length
        self.chunk_count = chunk_count if chunk_count is not None else len(self.chunks)
        self.pages = pages
        self.embedding_model = embedding_model
        self.embedding_dimension = embedding_dimension
        self.vector_store_status = vector_store_status
        self.error = error


def process_document(
    file_path: str,
    file_type: str,
    document_id: str = "doc_temp",
    filename: str = "doc",
    target_chunk_size: int = 1000,
    overlap: int = 150
) -> DocumentExtractionResult:
    """
    Complete document processing pipeline:
    1. Text extraction (PDF / DOCX / TXT) -> raw_text
    2. Text cleaning & normalization -> cleaned_text
    3. Document chunking -> List[DocumentChunk]
    4. Text embedding -> Embedded DocumentChunks
    5. Vector Storage -> Local Persistent ChromaDB
    Returns DocumentExtractionResult with complete pipeline output.
    """
    ext = file_type.lower().strip(".")
    
    # Step 1: Text Extraction
    if ext == "pdf":
        raw_res = _extract_pdf(file_path)
    elif ext == "docx":
        raw_res = _extract_docx(file_path)
    elif ext == "txt":
        raw_res = _extract_txt(file_path)
    else:
        return DocumentExtractionResult(
            success=False,
            file_type=ext,
            error=f"Unsupported file type extension: '.{ext}'"
        )

    if not raw_res.get("success"):
        return DocumentExtractionResult(
            success=False,
            file_type=ext,
            pages=raw_res.get("pages"),
            error=raw_res.get("error", "Failed to extract text from document.")
        )

    raw_text = raw_res.get("raw_text", "")
    pages = raw_res.get("pages")
    
    # Resource limit check: prevent text bomb exhaustion
    if len(raw_text) > MAX_DOCUMENT_CHARS:
        return DocumentExtractionResult(
            success=False,
            file_type=ext,
            pages=pages,
            error=f"Document exceeds maximum extractable character limit ({MAX_DOCUMENT_CHARS:,} characters)."
        )

    # Step 2: Text Cleaning & Normalization
    cleaned = clean_text(raw_text)

    if not cleaned:
        return DocumentExtractionResult(
            success=False,
            raw_text=raw_text,
            cleaned_text="",
            chunks=[],
            file_type=ext,
            text_length=len(raw_text),
            cleaned_text_length=0,
            chunk_count=0,
            pages=pages,
            error="Document contains no usable text after cleaning."
        )

    # Step 3: Document Chunking
    doc_chunks = chunk_text(
        cleaned_text=cleaned,
        document_id=document_id,
        filename=filename,
        target_chunk_size=target_chunk_size,
        overlap=overlap,
        file_type=ext,
        pages=pages
    )

    # Step 4: Text Embedding Generation
    embedder = EmbeddingService.get_instance()
    embedded_chunks = embedder.embed_chunks(doc_chunks)

    # Step 5: Vector Storage in Local ChromaDB
    try:
        vector_store = VectorStoreService.get_instance()
        # Clean old vectors if re-indexing document
        vector_store.delete_document_chunks(document_id)
        # Store new chunk vectors
        vector_store.add_chunks(embedded_chunks)
        store_status = "indexed"
    except Exception as e:
        return DocumentExtractionResult(
            success=False,
            raw_text=raw_text,
            cleaned_text=cleaned,
            chunks=embedded_chunks,
            file_type=ext,
            text_length=len(raw_text),
            cleaned_text_length=len(cleaned),
            chunk_count=len(embedded_chunks),
            pages=pages,
            vector_store_status="failed",
            error=f"Failed to persist vectors in ChromaDB: {str(e)}"
        )

    return DocumentExtractionResult(
        success=True,
        raw_text=raw_text,
        cleaned_text=cleaned,
        chunks=embedded_chunks,
        file_type=ext,
        text_length=len(raw_text),
        cleaned_text_length=len(cleaned),
        chunk_count=len(embedded_chunks),
        pages=pages,
        embedding_model=getattr(embedder, "active_model_name", embedder.model_name),
        embedding_dimension=embedder.dimension,
        vector_store_status=store_status
    )


def _extract_pdf(file_path: str) -> Dict[str, Any]:
    try:
        reader = PdfReader(file_path)
        num_pages = len(reader.pages)
        text_content = []
        
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text_content.append(page_text)
                
        full_text = "\n\n".join(text_content).strip()
        
        if not full_text:
            return {
                "success": False,
                "raw_text": "",
                "pages": num_pages,
                "error": "PDF file contains no extractable text (it may be a scanned image or empty PDF)."
            }
            
        return {
            "success": True,
            "raw_text": full_text,
            "pages": num_pages
        }
    except Exception as e:
        return {
            "success": False,
            "raw_text": "",
            "pages": None,
            "error": "Failed to process PDF document: file may be corrupted, password-protected, or malformed."
        }


def _extract_docx(file_path: str) -> Dict[str, Any]:
    try:
        doc = docx.Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
                    
        full_text = "\n".join(paragraphs).strip()
        
        if not full_text:
            return {
                "success": False,
                "raw_text": "",
                "pages": None,
                "error": "DOCX document contains no extractable text."
            }
            
        return {
            "success": True,
            "raw_text": full_text,
            "pages": None
        }
    except Exception as e:
        return {
            "success": False,
            "raw_text": "",
            "pages": None,
            "error": "Failed to process DOCX document: file may be corrupted or malformed."
        }


def _extract_txt(file_path: str) -> Dict[str, Any]:
    try:
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                full_text = f.read()
        except UnicodeDecodeError:
            with open(file_path, "r", encoding="latin-1") as f:
                full_text = f.read()
                
        full_text = full_text.strip()
        
        if not full_text:
            return {
                "success": False,
                "raw_text": "",
                "pages": None,
                "error": "TXT document contains no extractable text."
            }
            
        return {
            "success": True,
            "raw_text": full_text,
            "pages": None
        }
    except Exception as e:
        return {
            "success": False,
            "raw_text": "",
            "pages": None,
            "error": "Failed to process TXT document: file may be corrupted or in an unsupported encoding."
        }
